#!/usr/bin/env python3
# Build the speaker script for Milana as a Word document.
#
# The deck itself only carries what the audience sees; everything the presenter
# needs to say lives here. This is the same content that used to sit on the two
# "advice" slides, rewritten against the current ten-slide structure.
#
# Usage: python scripts/build-speaker-docx.py
#
# Content lives in slides/speaker-script.json so that the text can be reviewed
# as data rather than buried in formatting code.
import json
import pathlib

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
CONFIG = ROOT / "slides" / "speaker-script.json"
OUT = ROOT / "docs" / "Лапа рядом — текст выступления.docx"

# Palette and type shared with the deck (slides/lib/brand.js) so the script and
# the slides read as one project.
INK = RGBColor(0x17, 0x16, 0x1C)
VIOLET = RGBColor(0x6B, 0x2E, 0x8F)
ACCENT = RGBColor(0x9C, 0x25, 0x9B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEADER_FILL = "9C259B"   # magenta pill colour, white text on it
ZEBRA_FILL = "F2EAFB"    # the deck's lilac card tint

DISPLAY_FONT = "PT Sans Narrow"
BODY_FONT = "PT Sans"


def shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def set_cell_borders(table, hex_color="E5DEF1", size=4):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:color"), hex_color)
        borders.append(e)
    table._tbl.tblPr.append(borders)


def run_font(run, name, size, bold=False, italic=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    # Cyrillic must be routed through the same face, not the default.
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run_font(run, BODY_FONT, 9, color=VIOLET)
    for kind, instr in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def set_col_widths(table, widths_cm):
    """python-docx only honours a width set on every cell of the column, and
    only when autofit is off. Widths come from the JSON in centimetres."""
    table.autofit = False
    for i, cm_width in enumerate(widths_cm):
        width = Cm(cm_width)
        table.columns[i].width = width
        for cell in table.columns[i].cells:
            cell.width = width


def build_table(doc, block):
    headers = block.get("headers", [])
    rows = block.get("rows", [])
    table = doc.add_table(rows=0, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(table)

    # One header row, then one row per data row. Adding a row per cell would
    # stagger the header diagonally down the table.
    header = table.add_row()
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.text = ""
        shade(cell, HEADER_FILL)
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.space_before = Pt(2)
        run_font(para.add_run(text), BODY_FONT, 9.5, bold=True, color=WHITE)

    for r, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            if r % 2 == 1:
                shade(cells[i], ZEBRA_FILL)
            para = cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            run_font(para.add_run(text), BODY_FONT, 9.5, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    if block.get("widths"):
        set_col_widths(table, block["widths"])


def build(doc):
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    add_page_number_footer(section)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    run_font(title.add_run(cfg["title"]), DISPLAY_FONT, 26, bold=True, color=INK)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(2)
    run_font(sub.add_run(cfg["subtitle"]), BODY_FONT, 12.5, color=VIOLET)

    who = doc.add_paragraph()
    who.paragraph_format.space_after = Pt(14)
    run_font(who.add_run(cfg["author"]), BODY_FONT, 10.5, color=VIOLET)

    for sec in cfg["sections"]:
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(16)
        head.paragraph_format.space_after = Pt(6)
        head.paragraph_format.keep_with_next = True
        run_font(head.add_run(sec["heading"]), DISPLAY_FONT, 17, bold=True, color=ACCENT)

        for block in sec.get("content", []):
            kind = block.get("type")
            if kind == "paragraph":
                p = doc.add_paragraph()
                run_font(p.add_run(block["text"]), BODY_FONT, 11.5, color=INK)
            elif kind == "list":
                for item in block.get("items", []):
                    p = doc.add_paragraph(style="List Bullet")
                    p.paragraph_format.space_after = Pt(3)
                    run_font(p.add_run(item), BODY_FONT, 11.5, color=INK)
            elif kind == "table":
                build_table(doc, block)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    # The Windows console here is cp1252 and cannot print the Cyrillic filename.
    print(f"WROTE {OUT.name.encode('ascii', 'replace').decode('ascii')}")


if __name__ == "__main__":
    build(Document())